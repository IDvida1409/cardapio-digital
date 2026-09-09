import hashlib
import json
import os
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from email.parser import BytesParser
from email.policy import default as email_policy
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from urllib import request as urllib_request
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook

try:
    import psycopg
    from psycopg.rows import dict_row
except Exception:  # pragma: no cover - optional outside production
    psycopg = None
    dict_row = None


PARSER_VERSION = "backend-ai-v1"
DEFAULT_MODEL = "gemini-2.5-flash-lite"
MAX_UPLOAD_BYTES = 15 * 1024 * 1024


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def normalize_text(value):
    import unicodedata

    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    text = text.replace("_", " ").replace("|", " ").lower()
    return " ".join(text.split())


def cell_to_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def merged_range_for(sheet, row_number, col_number):
    for merged in sheet.merged_cells.ranges:
        if (
            merged.min_row <= row_number <= merged.max_row
            and merged.min_col <= col_number <= merged.max_col
        ):
            return merged
    return None


def read_workbook(file_bytes, filename):
    workbook = load_workbook(BytesIO(file_bytes), data_only=True)
    sheets = []
    useful_cells = []

    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        rows = []
        for row in sheet.iter_rows():
            cells = []
            for cell in row:
                text = cell_to_text(cell.value)
                if not normalize_text(text):
                    continue

                merged = merged_range_for(sheet, cell.row, cell.column)
                if merged and (cell.row != merged.min_row or cell.column != merged.min_col):
                    continue

                start_col = merged.min_col if merged else cell.column
                end_col = merged.max_col if merged else cell.column
                start_row = merged.min_row if merged else cell.row
                end_row = merged.max_row if merged else cell.row
                record = {
                    "sheet": sheet.title,
                    "sheetIndex": sheet_index,
                    "address": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "startRow": start_row,
                    "endRow": end_row,
                    "startColumn": start_col,
                    "endColumn": end_col,
                    "rowSpan": end_row - start_row + 1,
                    "columnSpan": end_col - start_col + 1,
                    "text": text,
                    "normalizedText": normalize_text(text),
                }
                cells.append(record)
                useful_cells.append(record)

            if cells:
                cells.sort(key=lambda item: (item["startColumn"], item["column"]))
                rows.append({"row": row[0].row, "cells": cells})

        sheets.append(
            {
                "name": sheet.title,
                "index": sheet_index,
                "rowCount": sheet.max_row,
                "columnCount": sheet.max_column,
                "rows": rows,
            }
        )

    return {
        "fileName": filename,
        "sha256": hashlib.sha256(file_bytes).hexdigest(),
        "sheetCount": len(sheets),
        "sheets": sheets,
        "usefulCells": useful_cells,
    }


def compact_workbook_for_ai(raw_workbook):
    blocks = []
    for sheet in raw_workbook["sheets"]:
        sheet_cells = []
        for row in sheet["rows"]:
            for cell in row["cells"]:
                sheet_cells.append(
                    {
                        "ref": f"{sheet['index']}!{cell['address']}",
                        "sheet": sheet["name"],
                        "row": cell["row"],
                        "col": cell["column"],
                        "rowSpan": cell["rowSpan"],
                        "colSpan": cell["columnSpan"],
                        "text": cell["text"],
                    }
                )
        blocks.append(
            {
                "sheet": sheet["name"],
                "index": sheet["index"],
                "rowCount": sheet["rowCount"],
                "columnCount": sheet["columnCount"],
                "cells": sheet_cells,
            }
        )
    return blocks


def expected_ai_schema():
    return {
        "type": "object",
        "properties": {
            "periodos": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "titulo": {"type": "string"},
                        "inicio": {"type": "string"},
                        "fim": {"type": "string"},
                    },
                    "required": ["titulo"],
                },
            },
            "dias": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "data": {"type": "string"},
                        "diaSemana": {"type": "string"},
                        "refeicoes": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "nome": {"type": "string"},
                                    "cardapios": {
                                        "type": "array",
                                        "items": {
                                            "type": "object",
                                            "properties": {
                                                "titulo": {"type": "string"},
                                                "tipos": {
                                                    "type": "array",
                                                    "items": {
                                                        "type": "object",
                                                        "properties": {
                                                            "titulo": {"type": "string"},
                                                            "itensComuns": {
                                                                "type": "array",
                                                                "items": {"type": "string"},
                                                            },
                                                            "sugestoes": {
                                                                "type": "array",
                                                                "items": {
                                                                    "type": "object",
                                                                    "properties": {
                                                                        "titulo": {"type": "string"},
                                                                        "itens": {
                                                                            "type": "array",
                                                                            "items": {"type": "string"},
                                                                        },
                                                                    },
                                                                    "required": ["titulo", "itens"],
                                                                },
                                                            },
                                                            "sourceRefs": {
                                                                "type": "array",
                                                                "items": {"type": "string"},
                                                            },
                                                        },
                                                        "required": ["titulo", "itensComuns", "sugestoes", "sourceRefs"],
                                                    },
                                                },
                                                "sourceRefs": {
                                                    "type": "array",
                                                    "items": {"type": "string"},
                                                },
                                            },
                                            "required": ["titulo", "tipos", "sourceRefs"],
                                        },
                                    },
                                },
                                "required": ["nome", "cardapios"],
                            },
                        },
                    },
                    "required": ["data", "refeicoes"],
                },
            },
            "celulasUsadas": {"type": "array", "items": {"type": "string"}},
            "celulasIgnoradas": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "ref": {"type": "string"},
                        "motivo": {"type": "string"},
                    },
                    "required": ["ref", "motivo"],
                },
            },
            "celulasPendentes": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "ref": {"type": "string"},
                        "motivo": {"type": "string"},
                    },
                    "required": ["ref", "motivo"],
                },
            },
            "confianca": {"type": "number"},
        },
        "required": ["dias", "celulasUsadas", "celulasIgnoradas", "celulasPendentes", "confianca"],
    }


def build_ai_prompt(raw_workbook):
    return {
        "instruction": (
            "Analise a grade de Excel como estrutura visual de um cardápio hospitalar. "
            "Não use conhecimento prévio sobre nomes específicos de dietas ou hospitais. "
            "Não invente dados. Preserve a hierarquia visual: cabeçalhos, blocos, subtítulos, sugestões e itens abaixo. "
            "Toda célula com texto útil deve estar em celulasUsadas, celulasIgnoradas com motivo, ou celulasPendentes. "
            "Se uma célula for nota operacional e não parte do cardápio, marque como ignorada com motivo. "
            "Responda somente JSON válido no schema solicitado."
        ),
        "schema": expected_ai_schema(),
        "workbook": {
            "fileName": raw_workbook["fileName"],
            "sha256": raw_workbook["sha256"],
            "sheets": compact_workbook_for_ai(raw_workbook),
        },
    }


def call_gemini(raw_workbook):
    api_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if not api_key:
        return None

    model = os.environ.get("GEMINI_MODEL", DEFAULT_MODEL).strip() or DEFAULT_MODEL
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    prompt = build_ai_prompt(raw_workbook)
    payload = {
        "contents": [
            {
                "role": "user",
                "parts": [{"text": json.dumps(prompt, ensure_ascii=False)}],
            }
        ],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json",
        },
    }
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib_request.Request(
        url,
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib_request.urlopen(req, timeout=90) as response:
            body = json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Gemini HTTP {exc.code}: {detail}") from exc
    except URLError as exc:
        raise RuntimeError(f"Erro de conexão com Gemini: {exc.reason}") from exc

    text = body["candidates"][0]["content"]["parts"][0]["text"]
    result = json.loads(text)
    result["_aiModel"] = model
    return result


def fallback_without_ai(raw_workbook):
    return {
        "dias": [],
        "celulasUsadas": [],
        "celulasIgnoradas": [],
        "celulasPendentes": [
            {
                "ref": f"{cell['sheetIndex']}!{cell['address']}",
                "motivo": "IA não configurada; célula aguardando classificação estrutural.",
            }
            for cell in raw_workbook["usefulCells"]
        ],
        "confianca": 0,
        "_aiModel": None,
    }


def validate_ai_result(raw_workbook, structured):
    all_refs = {f"{cell['sheetIndex']}!{cell['address']}" for cell in raw_workbook["usefulCells"]}
    used = set(structured.get("celulasUsadas") or [])
    ignored = {item.get("ref") for item in structured.get("celulasIgnoradas") or [] if item.get("ref")}
    pending = {item.get("ref") for item in structured.get("celulasPendentes") or [] if item.get("ref")}
    accounted = used | ignored | pending
    missing = sorted(all_refs - accounted)

    if missing:
        structured.setdefault("celulasPendentes", [])
        structured["celulasPendentes"].extend(
            {"ref": ref, "motivo": "Célula útil não classificada pelo modelo."}
            for ref in missing
        )
        pending.update(missing)

    return {
        "totalUsefulCells": len(all_refs),
        "usedCells": len(used),
        "ignoredCells": len(ignored),
        "pendingCells": len(pending),
        "missingCells": len(missing),
        "canPersist": len(pending) == 0 and float(structured.get("confianca") or 0) >= 0.85,
    }


class Store:
    def __init__(self):
        self.database_url = os.environ.get("DATABASE_URL", "").strip()
        self.sqlite_path = os.environ.get("SQLITE_PATH", "./nutrimenu.db")
        self.kind = "postgres" if self.database_url else "sqlite"
        self.ready = False
        self.init_error = ""

    def connect(self):
        if self.kind == "postgres":
            if psycopg is None:
                raise RuntimeError("psycopg não está instalado para conectar ao Postgres.")
            return psycopg.connect(self.database_url, row_factory=dict_row)
        return sqlite3.connect(self.sqlite_path)

    def init(self):
        with self.connect() as conn:
            cur = conn.cursor()
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS import_jobs (
                  id TEXT PRIMARY KEY,
                  filename TEXT NOT NULL,
                  file_hash TEXT NOT NULL,
                  status TEXT NOT NULL,
                  parser_version TEXT NOT NULL,
                  ai_model TEXT,
                  confidence REAL NOT NULL DEFAULT 0,
                  raw_cell_count INTEGER NOT NULL DEFAULT 0,
                  used_cell_count INTEGER NOT NULL DEFAULT 0,
                  ignored_cell_count INTEGER NOT NULL DEFAULT 0,
                  pending_cell_count INTEGER NOT NULL DEFAULT 0,
                  payload TEXT NOT NULL,
                  created_at TEXT NOT NULL
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS master_entities (
                  id TEXT PRIMARY KEY,
                  entity_type TEXT NOT NULL,
                  normalized_name TEXT NOT NULL,
                  display_name TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  UNIQUE(entity_type, normalized_name)
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS menu_days (
                  id TEXT PRIMARY KEY,
                  import_id TEXT NOT NULL,
                  date_text TEXT NOT NULL,
                  day_name TEXT,
                  source_sheet TEXT,
                  created_at TEXT NOT NULL
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS meal_cards (
                  id TEXT PRIMARY KEY,
                  day_id TEXT NOT NULL,
                  meal_name TEXT NOT NULL,
                  title TEXT NOT NULL,
                  order_index INTEGER NOT NULL DEFAULT 0,
                  source_refs TEXT NOT NULL,
                  created_at TEXT NOT NULL
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS menu_types (
                  id TEXT PRIMARY KEY,
                  card_id TEXT NOT NULL,
                  title TEXT NOT NULL,
                  order_index INTEGER NOT NULL DEFAULT 0,
                  source_refs TEXT NOT NULL,
                  created_at TEXT NOT NULL
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS menu_items (
                  id TEXT PRIMARY KEY,
                  type_id TEXT NOT NULL,
                  suggestion_title TEXT,
                  item_text TEXT NOT NULL,
                  source_refs TEXT NOT NULL,
                  created_at TEXT NOT NULL
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS cell_audit (
                  id TEXT PRIMARY KEY,
                  import_id TEXT NOT NULL,
                  source_ref TEXT NOT NULL,
                  sheet_name TEXT NOT NULL,
                  address TEXT NOT NULL,
                  text TEXT NOT NULL,
                  status TEXT NOT NULL,
                  reason TEXT,
                  created_at TEXT NOT NULL
                )
                """
            )
            conn.commit()
            self.ready = True
            self.init_error = ""

    def execute(self, conn, sql, values):
        if self.kind == "postgres":
            sql = sql.replace("?", "%s")
        conn.cursor().execute(sql, values)

    def save_import(self, raw_workbook, structured, validation):
        import_id = str(uuid.uuid4())
        status = "persisted" if validation["canPersist"] else "needs_review"
        timestamp = now_iso()

        with self.connect() as conn:
            self.execute(
                conn,
                """
                INSERT INTO import_jobs (
                  id, filename, file_hash, status, parser_version, ai_model, confidence,
                  raw_cell_count, used_cell_count, ignored_cell_count, pending_cell_count,
                  payload, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    import_id,
                    raw_workbook["fileName"],
                    raw_workbook["sha256"],
                    status,
                    PARSER_VERSION,
                    structured.get("_aiModel"),
                    float(structured.get("confianca") or 0),
                    validation["totalUsefulCells"],
                    validation["usedCells"],
                    validation["ignoredCells"],
                    validation["pendingCells"],
                    json.dumps(structured, ensure_ascii=False),
                    timestamp,
                ),
            )
            self.save_cell_audit(conn, import_id, raw_workbook, structured, timestamp)
            if validation["canPersist"]:
                self.save_menu(conn, import_id, structured, timestamp)
            conn.commit()

        return import_id, status

    def save_cell_audit(self, conn, import_id, raw_workbook, structured, timestamp):
        status_by_ref = {}
        for ref in structured.get("celulasUsadas") or []:
            status_by_ref[ref] = ("used", "")
        for item in structured.get("celulasIgnoradas") or []:
            status_by_ref[item.get("ref")] = ("ignored", item.get("motivo", ""))
        for item in structured.get("celulasPendentes") or []:
            status_by_ref[item.get("ref")] = ("pending", item.get("motivo", ""))

        for cell in raw_workbook["usefulCells"]:
            ref = f"{cell['sheetIndex']}!{cell['address']}"
            status, reason = status_by_ref.get(ref, ("pending", "Célula sem classificação."))
            self.execute(
                conn,
                """
                INSERT INTO cell_audit (
                  id, import_id, source_ref, sheet_name, address, text, status, reason, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    str(uuid.uuid4()),
                    import_id,
                    ref,
                    cell["sheet"],
                    cell["address"],
                    cell["text"],
                    status,
                    reason,
                    timestamp,
                ),
            )

    def save_menu(self, conn, import_id, structured, timestamp):
        for day in structured.get("dias") or []:
            day_id = str(uuid.uuid4())
            self.execute(
                conn,
                "INSERT INTO menu_days (id, import_id, date_text, day_name, source_sheet, created_at) VALUES (?, ?, ?, ?, ?, ?)",
                (day_id, import_id, day.get("data", ""), day.get("diaSemana", ""), "", timestamp),
            )
            for meal in day.get("refeicoes") or []:
                for card_index, card in enumerate(meal.get("cardapios") or []):
                    card_id = str(uuid.uuid4())
                    self.execute(
                        conn,
                        """
                        INSERT INTO meal_cards (
                          id, day_id, meal_name, title, order_index, source_refs, created_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            card_id,
                            day_id,
                            meal.get("nome", ""),
                            card.get("titulo", ""),
                            card_index,
                            json.dumps(card.get("sourceRefs") or [], ensure_ascii=False),
                            timestamp,
                        ),
                    )
                    for type_index, menu_type in enumerate(card.get("tipos") or []):
                        type_id = str(uuid.uuid4())
                        self.execute(
                            conn,
                            """
                            INSERT INTO menu_types (
                              id, card_id, title, order_index, source_refs, created_at
                            ) VALUES (?, ?, ?, ?, ?, ?)
                            """,
                            (
                                type_id,
                                card_id,
                                menu_type.get("titulo", ""),
                                type_index,
                                json.dumps(menu_type.get("sourceRefs") or [], ensure_ascii=False),
                                timestamp,
                            ),
                        )
                        for item in menu_type.get("itensComuns") or []:
                            self.save_item(conn, type_id, None, item, menu_type.get("sourceRefs") or [], timestamp)
                        for suggestion in menu_type.get("sugestoes") or []:
                            for item in suggestion.get("itens") or []:
                                self.save_item(conn, type_id, suggestion.get("titulo", ""), item, menu_type.get("sourceRefs") or [], timestamp)

    def save_item(self, conn, type_id, suggestion_title, item_text, source_refs, timestamp):
        clean = str(item_text or "").strip()
        if not clean:
            return
        self.execute(
            conn,
            "INSERT INTO menu_items (id, type_id, suggestion_title, item_text, source_refs, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (
                str(uuid.uuid4()),
                type_id,
                suggestion_title,
                clean,
                json.dumps(source_refs, ensure_ascii=False),
                timestamp,
            ),
        )
        entity_type = "preparacao"
        self.execute(
            conn,
            """
            INSERT INTO master_entities (id, entity_type, normalized_name, display_name, created_at)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT DO NOTHING
            """,
            (str(uuid.uuid4()), entity_type, normalize_text(clean), clean, timestamp),
        )


def parse_multipart(handler):
    content_type = handler.headers.get("Content-Type", "")
    if not content_type.startswith("multipart/form-data"):
        raise ValueError("Envie a planilha como multipart/form-data no campo 'file'.")

    length = int(handler.headers.get("Content-Length", "0"))
    if length <= 0 or length > MAX_UPLOAD_BYTES:
        raise ValueError("Arquivo vazio ou maior que o limite permitido.")

    body = handler.rfile.read(length)
    message = BytesParser(policy=email_policy).parsebytes(
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8")
        + body
    )
    if not message.is_multipart():
        raise ValueError("Corpo multipart inválido.")

    filename = ""
    file_bytes = b""
    for part in message.iter_parts():
        params = dict(part.get_params(header="content-disposition", failobj=[]) or [])
        if params.get("name") != "file":
            continue
        filename = os.path.basename(params.get("filename", ""))
        file_bytes = part.get_payload(decode=True) or b""
        break

    if not filename:
        raise ValueError("Campo 'file' não encontrado.")
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Envie um arquivo .xlsx.")
    if not file_bytes:
        raise ValueError("Arquivo vazio.")

    return filename, file_bytes


def allowed_origin(origin):
    allowed = [
        item.strip()
        for item in os.environ.get("ALLOWED_ORIGINS", "*").split(",")
        if item.strip()
    ]
    return "*" in allowed or origin in allowed


class Handler(BaseHTTPRequestHandler):
    store = Store()

    def end_headers(self):
        origin = self.headers.get("Origin", "")
        if allowed_origin(origin):
            self.send_header("Access-Control-Allow-Origin", origin or "*")
            self.send_header("Vary", "Origin")
        self.send_header("Access-Control-Allow-Methods", "GET,POST,OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type,Authorization")
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(HTTPStatus.NO_CONTENT)
        self.end_headers()

    def do_GET(self):
        parsed_path = urlparse(self.path)
        if parsed_path.path == "/health":
            self.write_json(
                {
                    "ok": self.store.ready,
                    "parserVersion": PARSER_VERSION,
                    "database": self.store.kind,
                    "databaseReady": self.store.ready,
                    "databaseError": self.store.init_error,
                    "geminiConfigured": bool(os.environ.get("GEMINI_API_KEY", "").strip()),
                }
            )
            return
        self.write_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)

    def do_POST(self):
        parsed_path = urlparse(self.path)
        if parsed_path.path != "/api/import-cardapio":
            self.write_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
            return

        try:
            query = parse_qs(parsed_path.query)
            include_result = query.get("includeResult") == ["1"]
            filename, file_bytes = parse_multipart(self)
            raw = read_workbook(file_bytes, filename)
            structured = call_gemini(raw) or fallback_without_ai(raw)
            validation = validate_ai_result(raw, structured)
            import_id, status = self.store.save_import(raw, structured, validation)
            payload = {
                "importId": import_id,
                "status": status,
                "parserVersion": PARSER_VERSION,
                "aiModel": structured.get("_aiModel"),
                "validation": validation,
                "summary": {
                    "fileName": raw["fileName"],
                    "sheetCount": raw["sheetCount"],
                    "fileHash": raw["sha256"],
                    "periodCount": len(structured.get("periodos") or []),
                    "dayCount": len(structured.get("dias") or []),
                },
            }
            if include_result:
                payload["result"] = structured

            self.write_json(payload, HTTPStatus.CREATED if status == "persisted" else HTTPStatus.ACCEPTED)
        except Exception as exc:
            self.write_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)

    def write_json(self, payload, status=HTTPStatus.OK):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def main():
    port = int(os.environ.get("PORT", "8000"))
    try:
        Handler.store.init()
    except Exception as exc:
        Handler.store.ready = False
        Handler.store.init_error = str(exc)
        print(f"Database initialization failed: {exc}", flush=True)
    server = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    print(f"NutriMenu AI API listening on :{port}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()

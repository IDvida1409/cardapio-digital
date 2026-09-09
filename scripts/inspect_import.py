import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_text(value):
    import unicodedata

    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return " ".join(text.replace("_", " ").replace("|", " ").lower().split())


def merged_range_for(sheet, row_number, col_number):
    for merged in sheet.merged_cells.ranges:
        if (
            merged.min_row <= row_number <= merged.max_row
            and merged.min_col <= col_number <= merged.max_col
        ):
            return merged
    return None


def extract_rows(sheet):
    rows = []
    for row in sheet.iter_rows():
        cells = []
        for cell in row:
            text = cell_text(cell.value)
            if not normalize_text(text):
                continue

            merged = merged_range_for(sheet, cell.row, cell.column)
            if merged and (cell.row != merged.min_row or cell.column != merged.min_col):
                continue

            start_col = merged.min_col if merged else cell.column
            end_col = merged.max_col if merged else cell.column
            key = f"{start_col}:{end_col}:{normalize_text(text)}"
            if any(current["key"] == key for current in cells):
                continue

            cells.append(
                {
                    "key": key,
                    "address": cell.coordinate,
                    "rowNumber": cell.row,
                    "colNumber": cell.column,
                    "startCol": start_col,
                    "endCol": end_col,
                    "text": text,
                }
            )

        if cells:
            cells.sort(key=lambda item: (item["startCol"], item["colNumber"]))
            rows.append(
                {
                    "number": row[0].row,
                    "cells": cells,
                    "text": " | ".join(cell["text"] for cell in cells),
                }
            )
    return rows


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: inspect_import.py input.xlsx output.json")

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    workbook = load_workbook(input_path, data_only=True)
    sheets = []

    for index, sheet in enumerate(workbook.worksheets, start=1):
        sheets.append(
            {
                "id": index,
                "name": sheet.title,
                "rowCount": sheet.max_row,
                "columnCount": sheet.max_column,
                "rows": extract_rows(sheet),
            }
        )

    output_path.write_text(
        json.dumps(
            {
                "fileName": input_path.name,
                "importedAt": "2026-09-09T00:00:00.000Z",
                "sheetCount": len(sheets),
                "sheets": sheets,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
